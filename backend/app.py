from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
import pdfplumber
import docx
from contradiction_checker import find_conflicts
from billing_tracker import billing
from external_monitor import get_external_policy
from fpdf import FPDF

app = FastAPI()

# Allow React frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def read_file(file: UploadFile):
    if file.filename.endswith(".pdf"):
        with pdfplumber.open(file.file) as pdf:
            return " ".join(page.extract_text() for page in pdf.pages)
    elif file.filename.endswith(".docx"):
        d = docx.Document(file.file)
        return " ".join(p.text for p in d.paragraphs)
    else:
        return file.file.read().decode("utf-8")

@app.post("/analyze/")
async def analyze(files: list[UploadFile] = File(...)):
    docs_text = [read_file(f) for f in files]
    billing.add_docs(len(files))
    conflicts = find_conflicts(docs_text)

    # Generate PDF report
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Doc Checker Agent Report", ln=True, align='C')
    for c in conflicts:
        pdf.multi_cell(0, 10, txt=c["message"])
    report_file = "report.pdf"
    pdf.output(report_file)

    billing.add_report()

    return {
        "conflicts": conflicts,
        "usage": {
            "docs_analyzed": billing.docs_analyzed,
            "reports_generated": billing.reports_generated,
            "bill": billing.total_bill()
        },
        "external_update": get_external_policy()
    }
